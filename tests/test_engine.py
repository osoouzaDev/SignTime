from __future__ import annotations

import io
import tempfile
import unittest
import zipfile
from pathlib import Path

import pandas as pd
from docx import Document

from contratos.engine import (
    analyze_inputs,
    format_cnpj,
    format_cpf,
    generate_contracts_zip,
)


class ContractEngineTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)
        self.template = self.directory / "modelo.docx"
        self.spreadsheet = self.directory / "dados.xlsx"

        document = Document()
        document.add_paragraph(
            "Cliente {{ nome_responsavel }} — CPF {{ cpf }} — {{ texto_endereco }} — {{ data }}"
        )
        document.save(self.template)

        pd.DataFrame(
            [
                {
                    "Seu nome completo:": "Ana da Silva",
                    "CPF": 123456789,
                    "Endereço Comercial Completo (Rua, Número, Bairro, Cidade/UF, CEP)": "Rua A, 10",
                }
            ]
        ).to_excel(self.spreadsheet, index=False)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_document_formatters_preserve_leading_zeroes(self) -> None:
        self.assertEqual(format_cpf(123456789), "001.234.567-89")
        self.assertEqual(format_cnpj(1234567800019), "01.234.567/8000-19")

    def test_analyze_suggests_existing_aliases_and_system_date(self) -> None:
        analysis = analyze_inputs(self.template, self.spreadsheet)
        self.assertEqual(analysis["rows"], 1)
        self.assertEqual(analysis["mapping"]["nome_responsavel"]["kind"], "column")
        self.assertEqual(analysis["mapping"]["cpf"]["value"], "CPF")
        self.assertEqual(analysis["mapping"]["texto_endereco"]["kind"], "address_clause")
        self.assertEqual(analysis["mapping"]["data"]["value"], "today")

    def test_generate_zip_contains_contracts_and_report(self) -> None:
        analysis = analyze_inputs(self.template, self.spreadsheet)
        output = self.directory / "contratos.zip"

        def fake_converter(docx_path: Path, pdf_path: Path) -> None:
            pdf_path.write_bytes(b"%PDF-1.4\n% test")

        summary = generate_contracts_zip(
            self.template,
            self.spreadsheet,
            analysis["mapping"],
            output,
            output_format="both",
            converter=fake_converter,
        )

        self.assertEqual(summary["generated"], 1)
        self.assertEqual(summary["errors"], 0)
        with zipfile.ZipFile(output) as archive:
            names = archive.namelist()
            self.assertIn("Ana da Silva_CPF.pdf", names)
            self.assertIn("Ana da Silva_CPF.docx", names)
            self.assertIn("relatorio_geracao.csv", names)

            generated_docx = Document(io.BytesIO(archive.read("Ana da Silva_CPF.docx")))
            text = "\n".join(paragraph.text for paragraph in generated_docx.paragraphs)
            self.assertIn("Ana da Silva", text)
            self.assertIn("001.234.567-89", text)
            self.assertIn("com endereço à Rua A, 10,", text)
            self.assertNotIn("{{", text)

    def test_pdf_failure_falls_back_to_docx(self) -> None:
        analysis = analyze_inputs(self.template, self.spreadsheet)
        output = self.directory / "fallback.zip"

        def failing_converter(docx_path: Path, pdf_path: Path) -> None:
            raise RuntimeError("conversor indisponível")

        summary = generate_contracts_zip(
            self.template,
            self.spreadsheet,
            analysis["mapping"],
            output,
            output_format="pdf",
            converter=failing_converter,
        )

        self.assertEqual(summary["pdf_fallbacks"], 1)
        with zipfile.ZipFile(output) as archive:
            self.assertIn("Ana da Silva_CPF.docx", archive.namelist())


if __name__ == "__main__":
    unittest.main()
